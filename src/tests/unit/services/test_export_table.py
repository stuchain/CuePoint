#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Exporting rows of named columns (CLEAN-11).

``ExportService.export_table`` is the row source "Export review list" writes
through. What matters about it:

- **each format holds exactly the columns asked for, in order**, and a missing
  value is an empty cell, not the word ``None``;
- **a file is whole or absent**: written beside its destination and renamed, so
  a failure leaves an existing file exactly as it was and no temporary behind;
- **an existing file is replaced only when asked**.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest

from cuepoint.exceptions.cuepoint_exceptions import ExportError
from cuepoint.services import export_service as export_module
from cuepoint.services.export_service import ExportService

COLUMNS = ("track_id", "title", "bpm", "match_disputed", "candidate_url")
ROWS = [
    {
        "track_id": 1,
        "title": "Café del Mar",
        "bpm": 124.5,
        "match_disputed": True,
        "candidate_url": "https://www.beatport.com/track/x/1",
        "not_a_column": "left out",
    },
    {
        "track_id": 2,
        "title": "Strobe",
        "bpm": None,
        "match_disputed": None,
        "candidate_url": None,
    },
]


@pytest.fixture
def service() -> ExportService:
    return ExportService()


def leftovers(folder: Path) -> list:
    return sorted(path.name for path in folder.iterdir() if path.suffix == ".tmp")


@pytest.mark.unit
class TestEachFormat:
    def test_csv_holds_the_columns_in_order_with_empty_cells_for_none(
        self, service, tmp_path
    ):
        target = tmp_path / "review.csv"
        service.export_table(COLUMNS, ROWS, str(target), "csv")

        raw = target.read_bytes()
        # A byte-order mark, so a spreadsheet opens "Café" as written.
        assert raw.startswith(b"\xef\xbb\xbf")
        with open(target, newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.reader(handle))
        assert rows[0] == list(COLUMNS)
        assert rows[1] == [
            "1",
            "Café del Mar",
            "124.5",
            "true",
            "https://www.beatport.com/track/x/1",
        ]
        assert rows[2] == ["2", "Strobe", "", "", ""]
        assert len(rows) == 3

    def test_json_is_objects_with_exactly_the_columns(self, service, tmp_path):
        target = tmp_path / "review.json"
        service.export_table(COLUMNS, ROWS, str(target), "json")

        data = json.loads(target.read_text(encoding="utf-8"))
        assert [list(row) for row in data] == [list(COLUMNS), list(COLUMNS)]
        assert data[0]["title"] == "Café del Mar"
        assert data[0]["match_disputed"] is True
        assert data[1]["bpm"] is None

    def test_excel_has_a_header_row_and_a_row_per_track(self, service, tmp_path):
        from openpyxl import load_workbook

        target = tmp_path / "review.xlsx"
        service.export_table(
            COLUMNS, ROWS, str(target), "excel", sheet_title="Review list"
        )

        workbook = load_workbook(target)
        sheet = workbook["Review list"]
        values = [[cell.value for cell in row] for row in sheet.iter_rows()]
        workbook.close()
        assert values[0] == list(COLUMNS)
        assert values[1][:3] == [1, "Café del Mar", 124.5]
        assert values[1][3] == "true"
        assert values[2][2] is None or values[2][2] == ""
        assert len(values) == 3

    def test_no_rows_is_a_header_alone(self, service, tmp_path):
        target = tmp_path / "empty.csv"
        service.export_table(COLUMNS, [], str(target), "csv")
        with open(target, newline="", encoding="utf-8-sig") as handle:
            assert list(csv.reader(handle)) == [list(COLUMNS)]


@pytest.mark.unit
class TestRefusals:
    def test_an_unknown_format(self, service, tmp_path):
        with pytest.raises(ExportError) as caught:
            service.export_table(COLUMNS, ROWS, str(tmp_path / "x.txt"), "txt")
        assert caught.value.error_code == "EXPORT_INVALID_FORMAT"
        assert not (tmp_path / "x.txt").exists()

    def test_an_existing_file_without_overwrite_is_left_alone(self, service, tmp_path):
        target = tmp_path / "review.csv"
        target.write_text("keep me", encoding="utf-8")

        with pytest.raises(ExportError) as caught:
            service.export_table(COLUMNS, ROWS, str(target), "csv")

        assert caught.value.error_code == "EXPORT_PATH_REFUSED"
        assert target.read_text(encoding="utf-8") == "keep me"

    def test_overwrite_replaces_it(self, service, tmp_path):
        target = tmp_path / "review.json"
        target.write_text("old", encoding="utf-8")

        service.export_table(COLUMNS, ROWS, str(target), "json", overwrite=True)

        assert json.loads(target.read_text(encoding="utf-8"))[1]["title"] == "Strobe"


@pytest.mark.unit
class TestWholeOrAbsent:
    def test_a_failed_write_leaves_the_existing_file_and_no_temporary(
        self, service, tmp_path, monkeypatch
    ):
        target = tmp_path / "review.csv"
        target.write_text("the previous export", encoding="utf-8")

        def fails(path, columns, rows):
            Path(path).write_text("half", encoding="utf-8")
            raise OSError("disk full")

        monkeypatch.setattr(export_module, "_write_csv_table", fails)

        with pytest.raises(ExportError) as caught:
            service.export_table(COLUMNS, ROWS, str(target), "csv", overwrite=True)

        assert caught.value.error_code == "EXPORT_WRITE_FAILED"
        assert "disk full" in caught.value.message
        assert target.read_text(encoding="utf-8") == "the previous export"
        assert leftovers(tmp_path) == []

    def test_a_successful_write_leaves_no_temporary(self, service, tmp_path):
        service.export_table(COLUMNS, ROWS, str(tmp_path / "a.csv"), "csv")
        assert leftovers(tmp_path) == []
